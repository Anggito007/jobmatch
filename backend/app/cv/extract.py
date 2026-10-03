"""Ekstraksi teks dari file CV (PDF / DOCX)."""
from io import BytesIO


def extract_text(content: bytes, filename: str) -> str:
    """Ekstrak teks mentah dari PDF atau DOCX berdasarkan ekstensi file.

    Kembalikan string kosong bila format tidak didukung atau parsing gagal.
    """
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            return _extract_pdf(content)
        if name.endswith(".docx"):
            return _extract_docx(content)
        if name.endswith(".txt") or name.endswith(".md"):
            return content.decode("utf-8", errors="ignore")
    except Exception:
        # Jangan jatuhkan seluruh request karena satu CV gagal diparse.
        return ""
    return ""


def _extract_pdf(content: bytes) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    return "\n".join(parts)


def _extract_docx(content: bytes) -> str:
    import docx

    doc = docx.Document(BytesIO(content))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    # Sertakan isi tabel (umum dipakai CV untuk riwayat/keahlian).
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
